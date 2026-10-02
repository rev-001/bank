"""
Unit & Integration Tests for Offline PDF-to-Excel Agent.
Verifies parsing accuracy, mathematical checks, template matching,
chat instruction learning, and Excel generation.
"""

import os
import tempfile
import unittest
from pathlib import Path
import openpyxl

from agent.core_agent import PDFToExcelAgent
from agent.memory import AgentMemory
from agent.parser import PDFParser
from agent.validator import ExactResultValidator
from agent.excel_exporter import ExcelExporter


class TestOfflineAgent(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db_path = "/home/user/history/test_agent_memory.db"
        cls.output_dir = "/home/user/outputs"
        cls.agent = PDFToExcelAgent(db_path=cls.db_path, default_output_dir=cls.output_dir)
        cls.samples_dir = "/home/user/samples"

    def test_01_standard_invoice_extraction(self):
        pdf_path = os.path.join(self.samples_dir, "invoice_standard.pdf")
        self.assertTrue(os.path.exists(pdf_path), "Sample PDF should exist")

        res = self.agent.process_pdf(pdf_path)
        
        self.assertIsNotNone(res["excel_path"])
        self.assertTrue(os.path.exists(res["excel_path"]))
        self.assertEqual(res["row_count"], 5)
        self.assertIn("Invoice Number", res["key_values"])
        self.assertEqual(res["key_values"]["Invoice Number"], "INV-2026-8842")
        self.assertGreaterEqual(res["accuracy_score"], 0.85)

    def test_02_multipage_po_extraction(self):
        pdf_path = os.path.join(self.samples_dir, "purchase_order_multipage.pdf")
        res = self.agent.process_pdf(pdf_path)
        
        # 8 items on page 1 + 4 items on page 2 = 12 items
        self.assertEqual(res["row_count"], 12)
        self.assertTrue(os.path.exists(res["excel_path"]))

    def test_03_chat_instruction_learning(self):
        # Teach the agent a new synonym mapping
        learn_res = self.agent.learn_from_chat("Rename 'Particulars' to 'Item Description'")
        self.assertEqual(learn_res["status"], "success")

        # Ingest noisy vendor bill
        pdf_path = os.path.join(self.samples_dir, "noisy_vendor_bill.pdf")
        res = self.agent.process_pdf(pdf_path)
        
        # Check that 'Particulars' was standardized to 'Item Description' or 'Description'
        self.assertTrue(any("Description" in h for h in res["headers"]))
        self.assertEqual(res["row_count"], 4)

    def test_04_mathematical_validation_exact_result(self):
        validator = ExactResultValidator()
        table = {
            "headers": ["Quantity", "Unit Price", "Total Price"],
            "rows": [
                [10, 50.0, 500.0],
                [5, 20.0, 100.0],
                [2, 150.0, ""]  # Missing total should be auto-calculated to 300.0
            ]
        }
        # Table sum is 500 + 100 + 300 = 900.0
        kv = {"Subtotal": "900.00", "Tax Amount": "90.00", "Grand Total": "990.00"}
        
        val_res = validator.validate_and_reconcile(table, kv)
        self.assertEqual(val_res["accuracy_score"], 1.0)
        self.assertEqual(val_res["table"]["rows"][2][2], 300.0)
        self.assertEqual(val_res["calculated_table_sum"], 900.0)

    def test_05_excel_formatting_and_formulas(self):
        out_file = "/home/user/outputs/test_verify.xlsx"
        exporter = ExcelExporter()
        
        table = {
            "headers": ["Item", "Quantity", "Unit Price", "Total Price"],
            "rows": [
                ["Item A", 10, 25.0, 250.0],
                ["Item B", 2, 100.0, 200.0]
            ],
            "column_types": ["text", "integer", "currency", "currency"]
        }
        kv = {"Invoice No": "TEST-101", "Grand Total": "$450.00"}
        
        exporter.export(out_file, table, kv)
        self.assertTrue(os.path.exists(out_file))

        # Inspect generated workbook with openpyxl
        wb = openpyxl.load_workbook(out_file, data_only=False)
        self.assertIn("Extracted Data", wb.sheetnames)
        self.assertIn("Document Summary", wb.sheetnames)
        self.assertIn("Agent Audit Log", wb.sheetnames)
        
        ws = wb["Extracted Data"]
        # Check if formula is present in total row
        has_formula = False
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=SUM"):
                    has_formula = True
                    break
        self.assertTrue(has_formula, "Excel file should include dynamic =SUM() formula")

    def test_06_bank_statement_transaction_layout_and_currency(self):
        headers = [
            "Txn Date",
            "Value Date",
            "Particulars",
            "Ref. No.",
            "Debit",
            "Credit",
            "Balance",
        ]
        table = {
            "headers": headers,
            "rows": [["01/08/26", "01/08/26", "Cash withdrawal", "ATM-001", 125.5, "", 45074.5]],
            "column_types": ["date", "date", "text", "text", "currency", "currency", "currency"],
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            out_file = os.path.join(temp_dir, "bank_statement.xlsx")
            ExcelExporter().export(
                out_file,
                table,
                {},
                metadata={"is_bank_statement": True, "currency": "INR"},
            )

            workbook = openpyxl.load_workbook(out_file, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Transactions"])
            worksheet = workbook["Transactions"]
            self.assertEqual(
                [cell.value for cell in worksheet[1]],
                headers,
            )
            self.assertEqual(worksheet.max_row, 2, "Transaction tab should contain no added totals")
            self.assertEqual(worksheet.max_column, 7)
            for cell_ref in ("E2", "G2"):
                self.assertEqual(worksheet[cell_ref].number_format, '"₹"#,##0.00')
            self.assertEqual(worksheet["E2"].value, 125.5)
            self.assertEqual(worksheet.freeze_panes, "A2")

    def test_07_bank_statement_parser_column_mapping(self):
        pdf_path = Path(__file__).resolve().parent / "samples" / "bank_statement.pdf"
        parsed = PDFParser().parse(str(pdf_path))

        self.assertTrue(parsed["is_bank_statement"])
        self.assertEqual(
            parsed["tables"][0]["headers"],
            ["Txn Date", "Value Date", "Particulars", "Ref. No.", "Debit", "Credit", "Balance"],
        )
        self.assertEqual(
            parsed["tables"][0]["rows"][0],
            ["2026-08-01", "", "Opening Balance", "INIT-001", "", "", "$45,200.00"],
        )


if __name__ == "__main__":
    unittest.main()
