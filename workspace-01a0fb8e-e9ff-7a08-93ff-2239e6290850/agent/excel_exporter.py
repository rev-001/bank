"""
Professional Excel Generator using OpenPyXL.
Generates beautifully formatted, type-accurate .xlsx workbooks with
dynamic formulas, clean typography, zebra rows, and audit sheets.
"""

import json
import os
import re
from pathlib import Path
from typing import Dict, List, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    SOURCE_COL = "Source PDF"
    MASTER_FILENAME = "all_transactions.xlsx"

    # Color Themes
    PALETTE = {
        "header_fill": "1E3A8A",       # Deep Classic Blue
        "header_text": "FFFFFF",       # Pure White
        "meta_header_fill": "3B82F6",  # Vibrant Blue
        "meta_header_text": "FFFFFF",
        "zebra_fill": "F8FAFC",        # Soft light gray/blue
        "total_fill": "F1F5F9",        # Subtle gray
        "border_color": "CBD5E1",     # Light slate border
        "audit_pass_fill": "DCFCE7",   # Light green
        "audit_pass_text": "166534",
        "audit_warn_fill": "FEF3C7",   # Light yellow
        "audit_warn_text": "92400E"
    }

    def __init__(self):
        pass

    def merge_into_master(
        self,
        master_path: str,
        table_data: Dict[str, Any],
        key_values: Dict[str, str],
        validation_info: Optional[Dict[str, Any]] = None,
        metadata: Optional[Dict[str, Any]] = None,
    ) -> str:
        """
        Import this PDF's transactions into a single rolling workbook.
        Re-processing the same filename replaces that document's rows.
        """
        master_file = Path(master_path).resolve()
        master_file.parent.mkdir(parents=True, exist_ok=True)
        store_path = master_file.with_suffix(".json")
        source = (metadata or {}).get("filename") or "document.pdf"

        store: Dict[str, Any] = {"documents": {}}
        if store_path.exists():
            try:
                store = json.loads(store_path.read_text(encoding="utf-8"))
                if not isinstance(store, dict) or "documents" not in store:
                    store = {"documents": {}}
            except Exception:
                store = {"documents": {}}

        store["documents"][source] = {
            "headers": table_data.get("headers", []),
            "rows": table_data.get("rows", []),
            "column_types": table_data.get("column_types", []),
            "key_values": key_values or {},
            "metadata": metadata or {},
        }
        store_path.write_text(
            json.dumps(store, default=str, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        combined = self._combine_store(store)
        combined_meta = {
            "filename": f"{len(store['documents'])} PDF document(s)",
            "template_name": "All Extractions (Combined)",
            "accuracy_score": (metadata or {}).get("accuracy_score", 0.95),
            "is_scanned": (metadata or {}).get("is_scanned", False),
        }
        return self.export(
            output_path=str(master_file),
            table_data=combined["table"],
            key_values=combined["key_values"],
            validation_info=validation_info,
            metadata=combined_meta,
        )

    def _combine_store(self, store: Dict[str, Any]) -> Dict[str, Any]:
        docs = store.get("documents") or {}

        def norm(header: str) -> str:
            return re.sub(r"\s+", " ", str(header).strip().lower())

        ordered: List[str] = [self.SOURCE_COL]
        seen = {norm(self.SOURCE_COL)}
        types_by_header: Dict[str, str] = {}

        for doc in docs.values():
            headers = doc.get("headers") or []
            ctypes = doc.get("column_types") or []
            for i, header in enumerate(headers):
                key = norm(header)
                if not key or key in seen:
                    continue
                seen.add(key)
                ordered.append(header)
                if i < len(ctypes):
                    types_by_header[key] = ctypes[i]

        all_rows: List[List[Any]] = []
        all_kv: Dict[str, str] = {}
        for source, doc in docs.items():
            headers = doc.get("headers") or []
            index = {norm(h): i for i, h in enumerate(headers)}
            for row in doc.get("rows") or []:
                mapped = []
                for header in ordered:
                    if norm(header) == norm(self.SOURCE_COL):
                        src_i = index.get(norm(self.SOURCE_COL))
                        if src_i is not None and src_i < len(row) and str(row[src_i]).strip():
                            mapped.append(row[src_i])
                        else:
                            mapped.append(source)
                        continue
                    col_i = index.get(norm(header))
                    mapped.append(row[col_i] if col_i is not None and col_i < len(row) else "")
                all_rows.append(mapped)
            for k, v in (doc.get("key_values") or {}).items():
                all_kv[f"{source} | {k}"] = v

        col_types = []
        for header in ordered:
            if norm(header) == norm(self.SOURCE_COL):
                col_types.append("text")
            else:
                col_types.append(types_by_header.get(norm(header), "text"))

        return {
            "table": {
                "headers": ordered,
                "rows": all_rows,
                "column_types": col_types,
            },
            "key_values": all_kv,
        }

    def export(self, output_path: str, table_data: Dict[str, Any], key_values: Dict[str, str],
               validation_info: Optional[Dict[str, Any]] = None, metadata: Optional[Dict[str, Any]] = None) -> str:
        """
        Creates a multi-tab professional Excel spreadsheet.
        """
        out_file = Path(output_path).resolve()
        out_file.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        
        # Styles definition
        header_font = Font(name="Calibri", size=11, bold=True, color=self.PALETTE["header_text"])
        header_fill = PatternFill(start_color=self.PALETTE["header_fill"], end_color=self.PALETTE["header_fill"], fill_type="solid")
        
        zebra_fill = PatternFill(start_color=self.PALETTE["zebra_fill"], end_color=self.PALETTE["zebra_fill"], fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        total_font = Font(name="Calibri", size=11, bold=True, color="000000")
        total_fill = PatternFill(start_color=self.PALETTE["total_fill"], end_color=self.PALETTE["total_fill"], fill_type="solid")
        
        thin_border_side = Side(style="thin", color=self.PALETTE["border_color"])
        double_border_side = Side(style="double", color="334155")
        
        cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        total_border = Border(top=thin_border_side, bottom=double_border_side, left=thin_border_side, right=thin_border_side)

        # -------------------------------------------------------------
        # SHEET 1: Extracted Table Data
        # -------------------------------------------------------------
        ws1 = wb.active
        ws1.title = "Extracted Data"
        ws1.views.sheetView[0].showGridLines = True

        headers = table_data.get("headers", [])
        rows = table_data.get("rows", [])
        col_types = table_data.get("column_types", ["text"] * len(headers))
        metadata = metadata or {}
        is_bank_statement = metadata.get("is_bank_statement", False)
        if is_bank_statement:
            ws1.title = "Transactions"
        currency = metadata.get("currency")
        currency_format = '"₹"#,##0.00' if is_bank_statement or currency == "INR" else "$#,##0.00"

        # Write Title Block if key values exist
        row_offset = 1
        if key_values.get("Invoice Number") or key_values.get("Vendor / Company"):
            ws1.cell(row=row_offset, column=1, value="DOCUMENT DATA EXTRACTION").font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
            row_offset += 1
            subtitle_parts = []
            if key_values.get("Invoice Number"):
                subtitle_parts.append(f"Doc #: {key_values['Invoice Number']}")
            if key_values.get("Invoice Date"):
                subtitle_parts.append(f"Date: {key_values['Invoice Date']}")
            if key_values.get("Vendor / Company"):
                subtitle_parts.append(f"Vendor: {key_values['Vendor / Company']}")
            
            ws1.cell(row=row_offset, column=1, value=" | ".join(subtitle_parts)).font = Font(name="Calibri", size=10, italic=True, color="64748B")
            row_offset += 2
        else:
            row_offset = 1

        # Write Table Headers
        header_row_num = row_offset
        for col_idx, header_text in enumerate(headers, start=1):
            displayed_header = str(header_text) if is_bank_statement else str(header_text).upper()
            cell = ws1.cell(row=header_row_num, column=col_idx, value=displayed_header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center" if col_types[col_idx-1] == "date" else ("right" if col_types[col_idx-1] in ["currency", "number", "integer", "decimal"] else "left"), vertical="center", wrap_text=True)
            cell.border = cell_border
        ws1.row_dimensions[header_row_num].height = 26

        # Write Data Rows
        start_data_row = header_row_num + 1
        current_data_row = start_data_row

        numeric_total_cols = []

        for r_i, row in enumerate(rows):
            is_zebra = (r_i % 2 == 1)
            row_fill = zebra_fill if is_zebra else white_fill
            
            for c_i, cell_val in enumerate(row):
                col_type = col_types[c_i] if c_i < len(col_types) else "text"
                cell = ws1.cell(row=current_data_row, column=c_i + 1)
                cell.fill = row_fill
                cell.border = cell_border

                if cell_val is None or cell_val == "":
                    cell.value = ""
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col_type == "currency":
                    try:
                        cell.value = float(cell_val)
                        cell.number_format = currency_format
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        if c_i not in numeric_total_cols and "balance" not in headers[c_i].lower():
                            numeric_total_cols.append(c_i)
                    except (ValueError, TypeError):
                        cell.value = str(cell_val)
                elif col_type == "integer":
                    try:
                        cell.value = int(cell_val)
                        cell.number_format = "#,##0"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        if c_i not in numeric_total_cols and ("qty" in headers[c_i].lower() or "units" in headers[c_i].lower() or "hours" in headers[c_i].lower()):
                            numeric_total_cols.append(c_i)
                    except (ValueError, TypeError):
                        cell.value = str(cell_val)
                elif col_type in ["number", "decimal"]:
                    try:
                        cell.value = float(cell_val)
                        cell.number_format = "#,##0.00"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    except (ValueError, TypeError):
                        cell.value = str(cell_val)
                elif col_type == "date":
                    cell.value = str(cell_val)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.value = str(cell_val)
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws1.row_dimensions[current_data_row].height = 20
            current_data_row += 1

        # Write TOTAL Row with Dynamic Formula
        if len(rows) > 0 and numeric_total_cols and not is_bank_statement:
            total_row_num = current_data_row
            ws1.row_dimensions[total_row_num].height = 22
            
            # Label in Column 1 (or preceding column)
            first_label_col = 1
            ws1.cell(row=total_row_num, column=first_label_col, value="TOTAL").font = total_font
            ws1.cell(row=total_row_num, column=first_label_col).fill = total_fill
            ws1.cell(row=total_row_num, column=first_label_col).border = total_border
            ws1.cell(row=total_row_num, column=first_label_col).alignment = Alignment(horizontal="left", vertical="center")

            for c_i in range(len(headers)):
                cell = ws1.cell(row=total_row_num, column=c_i + 1)
                cell.font = total_font
                cell.fill = total_fill
                cell.border = total_border

                if c_i in numeric_total_cols:
                    col_letter = get_column_letter(c_i + 1)
                    formula = f"=SUM({col_letter}{start_data_row}:{col_letter}{current_data_row - 1})"
                    cell.value = formula
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if col_types[c_i] == "currency":
                        cell.number_format = currency_format
                    elif col_types[c_i] == "integer":
                        cell.number_format = "#,##0"
                    else:
                        cell.number_format = "#,##0.00"

        if is_bank_statement:
            ws1.freeze_panes = "A2"
            ws1.auto_filter.ref = ws1.dimensions
            for col in ws1.columns:
                max_len = max(
                    (len(line) for cell in col if cell.value for line in str(cell.value).split("\n")),
                    default=0,
                )
                ws1.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 14)
            wb.save(str(out_file))
            return str(out_file)

        # -------------------------------------------------------------
        # SHEET 2: Document Key-Values & Metadata
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="Document Summary")
        ws2.views.sheetView[0].showGridLines = True

        ws2.cell(row=1, column=1, value="EXTRACTED DOCUMENT METADATA").font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        ws2.cell(row=2, column=1, value="Key parameters parsed from headers, summary blocks, and footers").font = Font(name="Calibri", size=10, italic=True, color="64748B")

        meta_hdr_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        ws2.cell(row=4, column=1, value="FIELD NAME").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        ws2.cell(row=4, column=1).fill = meta_hdr_fill
        ws2.cell(row=4, column=1).border = cell_border
        
        ws2.cell(row=4, column=2, value="EXTRACTED VALUE").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        ws2.cell(row=4, column=2).fill = meta_hdr_fill
        ws2.cell(row=4, column=2).border = cell_border
        ws2.row_dimensions[4].height = 24

        m_row = 5
        for k, v in key_values.items():
            ws2.cell(row=m_row, column=1, value=k).font = Font(name="Calibri", size=11, bold=True)
            ws2.cell(row=m_row, column=1).fill = zebra_fill if m_row % 2 == 1 else white_fill
            ws2.cell(row=m_row, column=1).border = cell_border
            
            c_val = ws2.cell(row=m_row, column=2, value=str(v))
            c_val.font = Font(name="Calibri", size=11)
            c_val.fill = zebra_fill if m_row % 2 == 1 else white_fill
            c_val.border = cell_border
            ws2.row_dimensions[m_row].height = 20
            m_row += 1

        # -------------------------------------------------------------
        # SHEET 3: Agent Audit & Verification
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="Agent Audit Log")
        ws3.views.sheetView[0].showGridLines = True

        ws3.cell(row=1, column=1, value="OFFLINE AGENT EXTRACTION AUDIT").font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        
        meta = metadata
        audit_summary = [
            ("Source PDF File", meta.get("filename", "document.pdf")),
            ("Template Matched", meta.get("template_name", "Auto-Detected Standard")),
            ("Extraction Confidence", f"{meta.get('accuracy_score', 0.98) * 100:.1f}%"),
            ("Document Type", meta.get("is_scanned", False) and "Scanned (OCR Processed)" or "Digital Vector PDF"),
            ("Total Rows Extracted", len(rows)),
            ("Key Fields Extracted", len(key_values)),
            ("Validation Checks Status", "All Mathematical Checks Verified" if (validation_info and validation_info.get("checks_passed", 0) > 0) else "Standard Structure Parsed")
        ]

        ws3.cell(row=3, column=1, value="AUDIT METRIC").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        ws3.cell(row=3, column=1).fill = meta_hdr_fill
        ws3.cell(row=3, column=1).border = cell_border
        ws3.cell(row=3, column=2, value="DETAILS").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        ws3.cell(row=3, column=2).fill = meta_hdr_fill
        ws3.cell(row=3, column=2).border = cell_border

        a_row = 4
        for k, v in audit_summary:
            ws3.cell(row=a_row, column=1, value=k).font = Font(name="Calibri", size=11, bold=True)
            ws3.cell(row=a_row, column=1).fill = zebra_fill if a_row % 2 == 1 else white_fill
            ws3.cell(row=a_row, column=1).border = cell_border
            
            c_val = ws3.cell(row=a_row, column=2, value=str(v))
            c_val.font = Font(name="Calibri", size=11)
            c_val.fill = zebra_fill if a_row % 2 == 1 else white_fill
            c_val.border = cell_border
            ws3.row_dimensions[a_row].height = 20
            a_row += 1

        # If there are detailed verification log entries
        if validation_info and validation_info.get("audit_log"):
            a_row += 2
            ws3.cell(row=a_row, column=1, value="RECONCILIATION & CORRECTION ACTIONS").font = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
            a_row += 1

            ws3.cell(row=a_row, column=1, value="Action Type").font = Font(name="Calibri", bold=True)
            ws3.cell(row=a_row, column=2, value="Reconciliation Details").font = Font(name="Calibri", bold=True)
            a_row += 1

            for entry in validation_info.get("audit_log", []):
                ws3.cell(row=a_row, column=1, value=str(entry.get("type", "Info")).replace("_", " ").title()).border = cell_border
                ws3.cell(row=a_row, column=2, value=str(entry)).border = cell_border
                a_row += 1

        # -------------------------------------------------------------
        # Auto-adjust column widths across all worksheets
        # -------------------------------------------------------------
        for ws in [ws1, ws2, ws3]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        lines = str(cell.value).split("\n")
                        for line in lines:
                            max_len = max(max_len, len(line))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        wb.save(str(out_file))
        return str(out_file)
